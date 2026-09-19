"""Shared test helpers."""

from __future__ import annotations

from datetime import datetime

from fastapi.testclient import TestClient

from app.db import InMemoryRepository
from app.db.base import Repository
from app.main import create_app
from app.schemas.common import EmploymentType, ExperienceLevel
from app.schemas.opportunities import NormalizedLocation, NormalizedOpportunity
from app.schemas.profiles import Profile


def make_opportunity(
    *,
    title: str = "Software Engineer",
    company: str = "Acme",
    city: str | None = "bengaluru",
    remote: bool = False,
    skills: list[str] | None = None,
    level: ExperienceLevel | None = ExperienceLevel.MID,
    india_eligible: str = "eligible",
    is_internship: bool = False,
    published: bool = True,
) -> NormalizedOpportunity:
    location = NormalizedLocation(
        raw="bengaluru, india" if city else "anywhere",
        city=city,
        state="karnataka" if city == "bengaluru" else None,
        country="india",
        workplace_type="remote" if remote else "on_site",
    )
    return NormalizedOpportunity(
        provider="test",
        external_id="x1",
        source_key="test",
        title=title,
        normalized_title=title.lower(),
        company_name=company,
        location=location,
        description_text="develop and ship software",
        salary_min=300000,
        salary_currency="INR",
        employment_type=EmploymentType.INTERNSHIP if is_internship else EmploymentType.FULL_TIME,
        experience_level=level,
        posted_at=datetime(2026, 9, 1),
        external_url=f"https://example.com/postings/{title}-{company}",
        apply_url=f"https://example.com/postings/{title}-{company}",
        india_eligible=india_eligible,
        is_internship=is_internship,
        published=published,
        skills=skills or [],
    )


def make_profile(user_id: str = "user-123", skills: list[str] | None = None) -> Profile:
    return Profile(
        user_id=user_id,
        full_name="Test Candidate",
        headline="Software Engineer",
        city="Bengaluru",
        state="Karnataka",
        skills=skills or ["python", "sql", "react"],
        experience=[
            {
                "role": "Software Engineer",
                "company": "Example Corp",
                "start_year": 2021,
                "end_year": 2026,
            }
        ],
        education=[{"degree": "B.Tech", "institution": "IIT", "field": "CSE"}],
    )


def seed_repository(repository: Repository, user_id: str = "user-123") -> InMemoryRepository:
    repository.upsert_opportunity(
        make_opportunity(title="Software Engineer", skills=["python", "sql", "react"]),
        "fp-eng",
    )
    repository.upsert_opportunity(
        make_opportunity(
            title="Frontend Developer",
            company="Femto",
            city="mumbai",
            skills=["react", "javascript"],
        ),
        "fp-fe",
    )
    repository.upsert_opportunity(
        make_opportunity(
            title="Remote Backend Intern",
            company="Startly",
            remote=True,
            skills=["python", "fastapi"],
            level=ExperienceLevel.FRESHER,
            is_internship=True,
        ),
        "fp-int",
    )
    return repository  # type: ignore[return-value]


def client_with_user(repository: Repository | None = None) -> TestClient:
    from app.api import deps
    from app.core.security import VerifiedUser, current_user

    repo = repository or seed_repository(InMemoryRepository())

    def override_repo():
        return repo

    async def override_current_user():
        return VerifiedUser(user_id="user-123", email="candidate@example.com")

    app = create_app()
    app.dependency_overrides[deps.get_repository] = override_repo
    app.dependency_overrides[current_user] = override_current_user
    return TestClient(app)


def make_docx_bytes(text: str | None = None) -> bytes:
    """Generate a small DOCX in memory with python-docx."""
    import io

    from docx import Document

    doc = Document()
    doc.add_heading("Test Candidate", level=0)
    doc.add_paragraph(text or "Software Engineer with 4 years of experience. Skills: Python, SQL, React.")
    doc.add_paragraph("Summary")
    doc.add_paragraph("Experience")
    doc.add_paragraph("Education")
    doc.add_paragraph("Skills")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def bearer(user_id: str = "user-123") -> dict[str, str]:
    return {"Authorization": "Bearer test-token"}
